"""Synthetic regressions for local feedback, never depending on private documents."""
import asyncio
from datetime import date
from io import BytesIO
from uuid import uuid4

import pytest
from docx import Document
from pydantic import TypeAdapter, ValidationError
from sqlmodel import select

from app.api import meeting as meeting_api
from app.api.beli import Adjustment, ProductInput
from app.api.changelog import ChangeLogCreate, ChangeLogUpdate
from app.api.iam import delete_job_title
from app.models.beli import BeliPolicy, BeliLedger, MorningAttendance
from app.models.iam import JobTitle, JobTitlePermission, Permission, User, UserStatus
from app.core.config import settings
from app.models.meeting import MeetingRecord, MeetingStatus, MeetingTranscriptSegment
from app.models.project import ProjectStage
from app.schemas import (JobTitleCreate, JobTitleUpdate, DepartmentCreate, DepartmentUpdate,
                         UserCreate, UserUpdate, BeliRuleCreate, BeliRuleUpdate)
from app.schemas.contract import CounterpartyCreate, ContractCreate, ContractUpdate
from app.schemas.finance import FinanceAccountCreate, FinanceAccountUpdate, TransactionCreate, TransactionUpdate
from app.schemas.project import ProjectCreate, ProjectUpdate, ProjectStageUpdate, StageStatusUpdate, ProjectTaskPlanUpdateRequest
from app.schemas.todo import TodoCreate, TodoUpdate, LeaveRequestCreate
from app.services.beli import reconcile_attendance, account_for
from test_beli_workflow import db, meeting, make_project  # Shared isolated SQLite fixture.
from test_meeting_transcript_import import _make_session


TEXT_FIELDS = [
    (ProjectCreate, 'name', True), (ProjectCreate, 'description', False),
    (ProjectUpdate, 'name', True), (ProjectUpdate, 'description', False),
    (TodoCreate, 'title', True), (TodoCreate, 'description', False), (TodoCreate, 'notes', False),
    (TodoUpdate, 'title', True), (TodoUpdate, 'description', False), (TodoUpdate, 'notes', False),
    (ProjectStageUpdate, 'deliverables', False),
    (StageStatusUpdate, 'deliverables', False),
    (ProjectTaskPlanUpdateRequest, 'title', True), (ProjectTaskPlanUpdateRequest, 'description', False),
    (ProductInput, 'name', True), (ProductInput, 'description', False), (LeaveRequestCreate, 'reason', False),
    (JobTitleCreate, 'name', True), (JobTitleCreate, 'description', False),
    (JobTitleUpdate, 'name', True), (JobTitleUpdate, 'description', False),
    (DepartmentCreate, 'name', True), (DepartmentCreate, 'description', False),
    (DepartmentUpdate, 'name', True), (DepartmentUpdate, 'description', False),
    (BeliRuleCreate, 'name', True), (BeliRuleCreate, 'note', False),
    (BeliRuleUpdate, 'name', True), (BeliRuleUpdate, 'note', False),
    (CounterpartyCreate, 'name', True), (CounterpartyCreate, 'address', False), (CounterpartyCreate, 'bank_name', False),
    (ChangeLogCreate, 'version', True), (ChangeLogCreate, 'title', True), (ChangeLogCreate, 'content', True),
    (ChangeLogUpdate, 'version', True), (ChangeLogUpdate, 'title', True), (ChangeLogUpdate, 'content', True),
    (ContractCreate, 'name', True), (ContractUpdate, 'name', True),
    (FinanceAccountCreate, 'account_name', True), (FinanceAccountCreate, 'bank_name', False), (FinanceAccountCreate, 'bank_branch', False),
    (FinanceAccountUpdate, 'account_name', True), (FinanceAccountUpdate, 'bank_name', False), (FinanceAccountUpdate, 'bank_branch', False),
    (TransactionCreate, 'purpose', False), (TransactionUpdate, 'purpose', False),
    (UserCreate, 'display_name', True), (UserUpdate, 'display_name', True), (Adjustment, 'reason', True),
]


@pytest.mark.parametrize('schema,field,required', TEXT_FIELDS, ids=lambda value: getattr(value, '__name__', str(value)))
def test_each_feedback_text_field(schema, field, required):
    adapter = TypeAdapter(schema.model_fields[field].rebuild_annotation())
    for value in ['   ', '\t\r\n', '\u3000']:
        with pytest.raises(ValidationError):
            adapter.validate_python(value)
    assert adapter.validate_python('  中间 two spaces  保留\n第二行  ') == '中间 two spaces  保留\n第二行'
    if required:
        for value in ['', None]:
            with pytest.raises(ValidationError): adapter.validate_python(value)
    else:
        assert adapter.validate_python('') == ''
        if schema.model_fields[field].default is None:
            assert adapter.validate_python(None) is None


@pytest.mark.parametrize('schema', [ProjectUpdate, TodoUpdate, JobTitleUpdate, DepartmentUpdate,
                                   BeliRuleUpdate, ContractUpdate, FinanceAccountUpdate, TransactionUpdate, UserUpdate, ChangeLogUpdate,
                                   ProjectStageUpdate, StageStatusUpdate])
def test_patch_omission_stays_omitted(schema):
    assert schema().model_dump(exclude_unset=True) == {}


def docx_bytes(text):
    buffer = BytesIO()
    document = Document()
    for line in text.splitlines(): document.add_paragraph(line)
    document.save(buffer)
    return buffer.getvalue()


def test_text_api_roundtrip_and_required_fields(db):
    session, boss, employee, other, entity, actor, client = db
    project = make_project(db)
    paths = [(f'/api/v1/project/projects/{project.id}', {'name': '  项目 a b  ', 'description': '  描述 a b  '}),
             (f'/api/v1/iam/users/{employee.id}', {'display_name': '  Employee a b  '})]
    for path, payload in paths:
        assert client.patch(path, json=payload).status_code == 200
    assert project.name == '项目 a b' and project.description == '描述 a b'
    assert employee.display_name == 'Employee a b'
    result = client.post('/api/v1/changelog/entries', json={'version': ' v1 ', 'title': ' 标题 ', 'content': ' 正文\n内部  空格 '})
    assert result.status_code == 200, result.text
    stored = client.get('/api/v1/changelog/entries').json()['data']['items'][0]
    assert stored['content'] == '正文\n内部  空格'
    for field in ['version', 'title', 'content']:
        assert client.patch(f"/api/v1/changelog/entries/{stored['id']}", json={field: '   '}).status_code == 422
    assert client.patch(f'/api/v1/iam/users/{employee.id}', json={'display_name': ''}).status_code == 422
    assert client.post('/api/v1/meeting/records', data={'title': '   '}).status_code in (400, 422)


@pytest.mark.parametrize('value', ['   ', '\t\r\n', '\u3000'])
def test_stage_deliverables_reject_blank_through_actual_patch_route(db, value):
    session, _, _, _, _, _, client = db
    project = make_project(db)
    stage = ProjectStage(project_id=project.id, stage_code='requirements', stage_name='需求',
                         sequence_no=1, deliverables='原有交付物')
    session.add(stage); session.commit()
    before_updated_at = stage.updated_at
    response = client.patch(f'/api/v1/project/projects/{project.id}/stages/{stage.id}',
                            json={'deliverables': value, 'planned_end_at': '2099-01-01'})
    assert response.status_code == 422, response.text
    assert response.json()['detail'][0]['loc'] == ['body', 'deliverables']
    session.refresh(stage)
    assert stage.deliverables == '原有交付物'
    assert stage.planned_end_at is None and stage.updated_at == before_updated_at


def test_stage_deliverables_normalize_clear_and_preserve_patch_fields(db):
    session, _, _, _, _, _, client = db
    project = make_project(db)
    stage = ProjectStage(project_id=project.id, stage_code='requirements', stage_name='需求',
                         sequence_no=1, deliverables='原有交付物')
    session.add(stage); session.commit()
    path = f'/api/v1/project/projects/{project.id}/stages/{stage.id}'
    expected = '交付物 two  spaces\n内部换行'
    cases = [({'deliverables': '  ' + expected + ' \t'}, expected),
             ({'planned_end_at': '2099-01-01'}, expected),
             ({'deliverables': None}, expected),
             ({'deliverables': ''}, ''),
             ({}, '')]
    for payload, text in cases:
        response = client.patch(path, json=payload)
        assert response.status_code == 200, response.text
        assert response.json()['data']['deliverables'] == text
        session.refresh(stage)
        assert stage.deliverables == text
        readback = client.get(f'/api/v1/project/projects/{project.id}/stages').json()['data']
        assert next(row for row in readback if row['id'] == str(stage.id))['deliverables'] == text
    # Serialized feature lists are not ordinary text and must not be normalized.
    feature_list = '  {"rows": []}  '
    assert client.patch(path, json={'feature_list': feature_list}).status_code == 200
    session.refresh(stage)
    assert stage.feature_list == feature_list


@pytest.mark.parametrize('value', ['-9999999999999', '999999999999999999', '0.003', 'NaN', 'Infinity'])
def test_money_rejected_without_side_effects(db, value):
    session, boss, employee, _, _, actor, client = db
    response = client.post(f'/api/v1/beli/accounts/{employee.id}/adjust', json={'amount': value, 'reason': '测试', 'request_key': str(uuid4())})
    assert response.status_code == 422, response.text
    assert not session.exec(select(BeliLedger)).all()
    assert account_for(session, employee.id)[1].balance == 2000
    response = client.post('/api/v1/beli/products', json={'name': '有效商品', 'kind': 'physical', 'price': value, 'stock': 1})
    assert response.status_code == 422


def test_adjustment_zero_and_boundaries(db):
    session, boss, employee, _, _, actor, client = db
    response = client.post(f'/api/v1/beli/accounts/{employee.id}/adjust', json={'amount': 0, 'reason': '测试', 'request_key': str(uuid4())})
    assert response.status_code == 400 and '不能为 0' in response.json()['message']
    for amount in ['9999999999.99', '-9999999999.99']:
        response = client.post(f'/api/v1/beli/accounts/{employee.id}/adjust', json={'amount': amount, 'reason': ' 边界 ', 'request_key': str(uuid4())})
        assert response.status_code == 200, response.text
    assert account_for(session, employee.id)[1].balance == 2000


def test_pagination_totals_and_legacy_arrays(db):
    session, boss, employee, _, _, actor, client = db
    for index in range(30):
        session.add(BeliLedger(user_id=employee.id, event_key=f'fixture:{index}', kind='manual', amount=1, balance_after=1, reason='测试'))
    session.commit()
    path = f'/api/v1/beli/ledger?user_id={employee.id}'
    assert len(client.get(path).json()['data']) == 30
    data = client.get(path + '&paginated=true').json()['data']
    assert data['total'] == 30 and len(data['items']) == 30
    assert client.get(path + '&paginated=true&page=2').json()['data']['items'] == []
    actor[0] = employee
    for resource in ['ledger', 'rewards', 'orders']:
        params = {'paginated': True, 'user_id': str(employee.id)}
        assert 'total' in client.get(f'/api/v1/beli/{resource}', params=params).json()['data']


@pytest.mark.parametrize('filename,content', [('bad.doc', b'bad'), ('bad.docx', b'broken'), ('empty.docx', b''), ('empty.docx', docx_bytes(''))])
def test_failed_meeting_import_leaves_no_record_or_charges(db, filename, content):
    session, boss, employee, _, _, actor, client = db
    result = client.post('/api/v1/meeting/records', data={'title': '无效导入', 'meeting_date': '2026-10-20'},
                         files={'transcript_file': (filename, content)})
    assert result.status_code == 400, result.text
    assert not session.exec(select(MeetingRecord)).all()
    assert not session.exec(select(MorningAttendance)).all()
    assert not session.exec(select(BeliLedger)).all()


def test_atomic_import_template_and_replacement_failure(db):
    session, boss, employee, _, _, actor, client = db
    result = client.post('/api/v1/meeting/records', data={'title': '  导入测试  ', 'meeting_type': 'other'},
                         files={'transcript_file': ('meeting.docx', docx_bytes("'employee'：正文有 O'Connor\n'外部'：测试"))})
    assert result.status_code == 200, result.text
    row = session.get(MeetingRecord, result.json()['data']['id'])
    assert row.title == '导入测试'
    assert row.speaker_mapping == {'speaker_1': 'employee', 'speaker_2': '外部'}
    segments = session.exec(select(MeetingTranscriptSegment).where(MeetingTranscriptSegment.meeting_id == row.id)).all()
    assert len(segments) == 2 and "O'Connor" in segments[0].content
    result = client.post(f'/api/v1/meeting/records/{row.id}/upload-transcript', files={'file': ('bad.pdf', b'bad')})
    assert result.status_code == 400
    assert session.get(MeetingRecord, row.id).speaker_mapping['speaker_1'] == 'employee'
    assert len(session.exec(select(MeetingTranscriptSegment)).all()) == 2
    response = client.get('/api/v1/meeting/transcript-template')
    assert response.status_code == 200
    assert 'application/vnd.openxmlformats' in response.headers['content-type']
    assert len(meeting_api._parse_transcript_text('\n'.join(p.text for p in Document(BytesIO(response.content)).paragraphs))[0]) == 3


def test_pdf_import_preserves_speakers(db):
    from PyPDF2 import PdfWriter
    from PyPDF2.generic import DictionaryObject, NameObject, DecodedStreamObject
    _, _, _, _, _, _, client = db
    writer = PdfWriter()
    writer.add_blank_page(width=300, height=300)
    page = writer.pages[0]
    font = DictionaryObject({NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/Type1'),
                             NameObject('/BaseFont'): NameObject('/Helvetica')})
    page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): font})})
    stream = DecodedStreamObject()
    stream.set_data(b'BT /F1 12 Tf 10 250 Td (employee: Good morning) Tj 0 -20 Td (other: Hello) Tj ET')
    page[NameObject('/Contents')] = stream
    buffer = BytesIO(); writer.write(buffer)
    response = client.post('/api/v1/meeting/records', data={'title': 'PDF', 'meeting_type': 'other'},
        files={'transcript_file': ('meeting.pdf', buffer.getvalue(), 'application/pdf')})
    assert response.status_code == 200, response.text
    saved = client.get(f"/api/v1/meeting/records/{response.json()['data']['id']}").json()['data']
    assert saved['speaker_mapping'] == {'speaker_1': 'employee', 'speaker_2': 'other'}


def test_meeting_import_settlement_failure_rolls_back(db, monkeypatch):
    session, boss, employee, _, _, actor, client = db
    def fail(*args): raise RuntimeError('synthetic settlement failure')
    monkeypatch.setattr(meeting_api, '_settle_attendance', fail)
    response = client.post('/api/v1/meeting/records', data={'title': '回滚测试'},
        files={'transcript_file': ('meeting.docx', docx_bytes('讲话人1：测试'))})
    assert response.status_code == 500
    session.rollback()
    assert not session.exec(select(MeetingRecord)).all()
    assert not session.exec(select(MeetingTranscriptSegment)).all()
    assert not session.exec(select(BeliLedger)).all()


def test_successful_asr_cannot_reuse_old_speaker_identity(db, monkeypatch):
    session, boss, employee, other, _, _, _ = db
    row = meeting(session, boss, date(2026, 10, 20), [employee.display_name])
    row.speaker_mapping = {'speaker_1': employee.display_name}
    row.speaker_user_mapping = {'speaker_1': str(employee.id)}
    session.add(row); session.commit()
    monkeypatch.setattr(meeting_api, 'engine', session.get_bind())
    async def synthetic_asr(*args):
        return {'segments': [{'speaker_id': 'speaker_1', 'content': '新说话人', 'start_time': 0, 'end_time': 1}]}
    monkeypatch.setattr(meeting_api, 'submit_asr_task', synthetic_asr)
    asyncio.run(meeting_api._run_asr(row.id))
    session.expire_all()
    assert row.speaker_mapping == {} and row.speaker_user_mapping == {}
    assert row.attendees == ['speaker_1']
    assert account_for(session, employee.id)[1].balance == 1500


def test_employee_id_binding_external_name_legacy_and_rename(db):
    session, boss, employee, other, _, actor, client = db
    employee.display_name = other.display_name = '同名'
    session.add_all([employee, other]); session.flush()
    day = date(2026, 10, 20)
    row = meeting(session, boss, day, ['同名'])
    row.speaker_mapping = {'speaker_1': '同名'}
    session.add(MeetingTranscriptSegment(meeting_id=row.id, speaker_id='speaker_1', content='发言'))
    session.add(row); session.commit()
    reconcile_attendance(session, day); session.commit()
    assert account_for(session, employee.id)[1].balance == 1500
    result = client.get('/api/v1/beli/attendance', params={'day': str(day)}).json()['data']
    assert result['unmatched_names'] == ['同名']
    payload = {'speaker_mapping': {'speaker_1': '同名', 'unused': other.username},
               'speaker_user_mapping': {'speaker_1': str(employee.id), 'unused': str(other.id)}}
    for _ in range(2):
        response = client.put(f'/api/v1/meeting/records/{row.id}/speakers', json=payload)
        assert response.status_code == 200, response.text
    assert account_for(session, employee.id)[1].balance == 2000
    assert account_for(session, other.id)[1].balance == -500  # Unused speaker never counts.
    assert client.get('/api/v1/beli/attendance', params={'day': str(day)}).json()['data']['unmatched_names'] == []
    employee.display_name = '改名'; session.add(employee); session.flush()
    reconcile_attendance(session, day); session.commit()
    assert account_for(session, employee.id)[1].balance == 2000
    # Old clients do not erase unchanged bindings.
    client.put(f'/api/v1/meeting/records/{row.id}/speakers', json={'speaker_mapping': row.speaker_mapping})
    assert row.speaker_user_mapping['speaker_1'] == str(employee.id)
    # An explicit external name must not accidentally match an employee username.
    response = client.put(f'/api/v1/meeting/records/{row.id}/speakers', json={
        'speaker_mapping': {'speaker_1': employee.username}, 'speaker_user_mapping': {'speaker_1': None}})
    assert response.status_code == 200
    assert account_for(session, employee.id)[1].balance == 1500
    # Unbound legacy labels resolve a unique username before a nonunique name.
    client.put(f'/api/v1/meeting/records/{row.id}/speakers', json={
        'speaker_mapping': {'speaker_1': employee.username}, 'speaker_user_mapping': {}})
    assert account_for(session, employee.id)[1].balance == 2000


def test_invalid_binding_directory_and_pre_activation(db):
    session, boss, employee, other, _, actor, client = db
    directory = client.get('/api/v1/meeting/employees').json()['data']
    assert set(directory[0]) == {'id', 'display_name', 'username'}
    assert str(employee.id) in {row['id'] for row in directory}
    other.status = UserStatus.INACTIVE; session.add(other); session.commit()
    assert str(other.id) not in {row['id'] for row in client.get('/api/v1/meeting/employees').json()['data']}
    row = meeting(session, boss, date(2026, 10, 20))
    for uid in [str(uuid4()), str(other.id)]:
        assert client.put(f'/api/v1/meeting/records/{row.id}/speakers', json={
            'speaker_mapping': {'speaker_1': '名字'}, 'speaker_user_mapping': {'speaker_1': uid}}).status_code == 400
    policy = session.exec(select(BeliPolicy)).one()
    policy.starts_on = date(2026, 10, 16); session.add(policy); session.commit()
    result = client.get('/api/v1/beli/attendance', params={'day': '2026-10-15'}).json()['data']
    assert result['policy_active'] is False and result['starts_on'] == '2026-10-16'
    assert result['items'] == []
    reconcile_attendance(session, date(2026, 10, 15)); session.commit()
    assert not session.exec(select(BeliLedger)).all()


def test_employee_directory_requires_only_meeting_permission(db, monkeypatch):
    session, boss, employee, other, _, actor, client = db
    actor[0] = employee
    monkeypatch.setattr(settings, 'ENFORCE_RBAC', True)
    assert client.get('/api/v1/meeting/employees').status_code == 403
    title = JobTitle(name='会议记录员')
    permission = Permission(code='meeting.write', name='会议编辑', module='meeting')
    session.add_all([title, permission]); session.flush()
    employee.job_title_id = title.id
    session.add_all([employee, JobTitlePermission(job_title_id=title.id, permission_id=permission.id)]); session.commit()
    assert client.get('/api/v1/meeting/employees').status_code == 200
    assert client.get('/api/v1/iam/users').status_code == 403


def test_in_use_job_title_cannot_be_deleted(db):
    session, boss, employee, _, _, _, client = db
    title = JobTitle(name='有人使用的职位')
    session.add(title); session.flush()
    employee.job_title_id = title.id; session.add(employee); session.commit()
    assert client.delete(f'/api/v1/iam/job-titles/{title.id}').status_code == 400
    assert session.get(JobTitle, title.id) is not None


def test_delete_job_title_cleans_permission_links_with_foreign_keys():
    with _make_session() as session:
        title = JobTitle(name='测试职位')
        permission = Permission(code='fixture.read', name='测试权限', module='fixture')
        session.add_all([title, permission]); session.flush()
        session.add(JobTitlePermission(job_title_id=title.id, permission_id=permission.id)); session.commit()
        asyncio.run(delete_job_title(title.id, session, current_user=None))
        assert session.get(JobTitle, title.id) is None
        assert session.get(Permission, permission.id) is not None
        assert not session.exec(select(JobTitlePermission)).all()
